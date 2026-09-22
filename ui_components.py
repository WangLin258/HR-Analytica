# -*- coding: utf-8 -*-
"""HR Analytica - Streamlit UI components."""

import streamlit as st
import pandas as pd
import json, base64, os, tempfile
import numpy as np
from pathlib import Path
from datetime import datetime
from typing import Optional, List, Union

from analysis_engine import (
    read_file, clean_data, detect_column_types, basic_stats, fmt_val,
    plot_bars, plot_hist, plot_rank, plot_corr, gen_summary,
    calc_penetration, check_internal_fairness, gen_salary_advice,
    recruit_metrics, recruit_funnel, recruit_summary, find_performance_columns,
    find_cost_column, find_channel_column, find_job_column, find_total_cost_column,
    recruit_cross_stats, plot_heatmap, plot_line,
    export_pdf, export_excel, CONFIG, CHART_COLORS,
    init_db, save_history, save_report_image, load_recent_history, get_history_images,
    cleanup_old_history, find_level_column, salary_bandwidth_analysis,
    plot_bandwidth_box, plot_compa_distribution, gen_bandwidth_advice,
    export_history_excel, export_history_pdf,
    recruit_cycle_metrics, plot_cycle_bar, gen_cycle_advice,
    data_quality_report, save_user_pref, load_user_pref,
)

from analysis_engine import ENCODINGS
from config import BASE_DIR


def load_css() -> None:
    """Load external style.css."""
    css_path = BASE_DIR / "style.css"
    if css_path.exists():
        st.markdown(f"<style>{css_path.read_text(encoding='utf-8')}</style>", unsafe_allow_html=True)


def stat_card(label: str, value: str, theme: str = "default") -> None:
    """Render a themed stat card."""
    st.markdown(
        f'<div class="stat-card {theme}"><div class="label">{label}</div>'
        f'<div class="value">{value}</div></div>',
        unsafe_allow_html=True,
    )


def chart_card(title: str, fig) -> None:
    """Render a chart inside a styled card."""
    st.markdown(f'<div class="chart-card"><div class="card-title">{title}</div></div>', unsafe_allow_html=True)
    st.plotly_chart(fig, use_container_width=True, config={"displayModeBar": False})


def csv_download_button(df: pd.DataFrame, filename: str, key: str, label: str = "下载 CSV") -> None:
    """Small CSV download button next to analysis tables."""
    data = df.to_csv(index=True).encode("utf-8-sig")
    st.download_button(label, data=data, file_name=filename, mime="text/csv", key=key)


def read_file_large(uploaded_file, sheet: Optional[str] = None):
    """Read CSV/Excel with chunked loading and progress feedback."""
    name = uploaded_file.name.lower()
    if name.endswith(".csv"):
        for enc in ENCODINGS:
            try:
                uploaded_file.seek(0)
                # Count newlines cheaply for progress estimate
                count = 0
                while True:
                    buf = uploaded_file.read(1024 * 1024)
                    if not buf:
                        break
                    count += buf.count(b"\n")
                uploaded_file.seek(0)
                reader = pd.read_csv(uploaded_file, encoding=enc, chunksize=CONFIG["chunk_rows"])
                chunks = []
                with st.status("正在读取文件...", expanded=True) as status:
                    for i, chunk in enumerate(reader):
                        chunks.append(chunk)
                        start = i * CONFIG["chunk_rows"] + 1
                        end = min((i + 1) * CONFIG["chunk_rows"], max(count, 1))
                        status.update(label=f"正在读取第 {start}-{end} 行...")
                        st.progress(min((i + 1) * CONFIG["chunk_rows"] / max(count, 1), 1.0))
                    status.update(label="文件读取完成", state="complete", expanded=False)
                df = pd.concat(chunks, ignore_index=True)
                return df, f"CSV ({enc})", []
            except (UnicodeDecodeError, UnicodeError):
                continue
        raise ValueError("无法识别 CSV 编码")
    else:
        uploaded_file.seek(0)
        if name.endswith(".xlsx"):
            from openpyxl import load_workbook
            wb = load_workbook(uploaded_file, read_only=True, data_only=True)
            sh = sheet or wb.sheetnames[0]
            ws = wb[sh]
            rows_iter = ws.iter_rows(values_only=True)
            header = next(rows_iter, None)
            if header is None:
                wb.close()
                return pd.DataFrame(), "Excel", wb.sheetnames
            total = max(ws.max_row - 1, 0)
            chunks = []
            chunk = []
            with st.status("正在读取Excel文件...", expanded=True) as status:
                for i, row in enumerate(rows_iter, start=1):
                    chunk.append(row)
                    if len(chunk) >= CONFIG["chunk_rows"]:
                        chunks.append(pd.DataFrame(chunk, columns=header))
                        start_row = max(i - CONFIG["chunk_rows"] + 1, 1)
                        status.update(label=f"正在读取第 {start_row}-{i} 行 / 共约 {total} 行...")
                        st.progress(min(i / max(total, 1), 1.0))
                        chunk = []
                if chunk:
                    chunks.append(pd.DataFrame(chunk, columns=header))
                status.update(label="文件读取完成", state="complete", expanded=False)
            sheetnames = wb.sheetnames
            wb.close()
            df = pd.concat(chunks, ignore_index=True) if chunks else pd.DataFrame(columns=header)
            return df, "Excel", sheetnames
        else:
            xl = pd.ExcelFile(uploaded_file)
            sh = sheet or xl.sheet_names[0]
            with st.status("正在读取Excel文件...", expanded=True) as status:
                df = pd.read_excel(xl, sheet_name=sh, engine="xlrd")
                status.update(label="文件读取完成", state="complete", expanded=False)
            return df, "Excel", xl.sheet_names


def render_mobile_hint() -> None:
    """Show mobile optimization note in sidebar."""
    with st.sidebar:
        st.markdown(
            '<p style="font-size:0.8rem;color:#94A3B8;text-align:center;margin-top:1rem;">'
            "本工具为桌面端优化，移动端体验正在完善中。</p>",
            unsafe_allow_html=True,
        )


def render_nav() -> None:
    """Sidebar navigation with active highlight."""
    items = [("首页", "home"), ("薪酬分析", "salary"), ("招聘分析", "recruit"), ("历史报告", "report")]
    icons = {"home": "▢", "salary": "◫", "recruit": "☰", "report": "⊙"}
    with st.sidebar:
        st.markdown('<div class="brand-title">HR Analytica</div>', unsafe_allow_html=True)
        for label, key in items:
            active = st.session_state.page == key
            icon = icons.get(key, "")
            if st.button(f"{icon} {label}", key=f"nav_{key}", use_container_width=True,
                         type="primary" if active else "secondary"):
                st.session_state.page = key
                st.rerun()
        st.markdown("<br><br>", unsafe_allow_html=True)
        with st.expander("高级设置"):
            frac_pct = st.slider(
                "抽样比例",
                min_value=10,
                max_value=100,
                value=int(CONFIG["sample_frac"] * 100),
                step=5,
                key="sample_frac_slider",
            )
            st.session_state.sample_frac = frac_pct / 100
            st.caption("抽样比例越低，分析速度越快，但结果代表性越弱。")
        with st.expander("API"):
            api_enabled = st.toggle("启用 API", value=False, key="api_toggle")
            st.session_state.api_enabled = api_enabled
            if api_enabled:
                st.caption("POST /api/analyze")
                st.caption("运行: uvicorn api:app --port 8000")
        with st.expander("外观"):
            dark = st.toggle("暗色模式",
                             value=load_user_pref("dark_mode", "false") == "true",
                             key="dark_toggle")
            st.session_state.dark_mode = dark
            save_user_pref("dark_mode", "true" if dark else "false")
        with st.expander("报告模板"):
            company = st.text_input("公司名称", value=load_user_pref("company_name", "HR Analytica"),
                                    key="tpl_company")
            title = st.text_input("报告标题", value=load_user_pref("report_title", "人力资源数据分析报告"),
                                  key="tpl_title")
            reporter = st.text_input("报告人", value=load_user_pref("reporter", "HR Analytica 自动生成"),
                                     key="tpl_reporter")
            save_user_pref("company_name", company)
            save_user_pref("report_title", title)
            save_user_pref("reporter", reporter)
            logo = st.file_uploader("Logo", type=["png", "jpg", "jpeg"], key="tpl_logo")
            if logo:
                logo_b64 = base64.b64encode(logo.getvalue()).decode()
                save_user_pref("logo_b64", logo_b64)
                st.session_state.saved_logo = logo_b64
            elif "saved_logo" not in st.session_state:
                st.session_state.saved_logo = load_user_pref("logo_b64", "")
        st.caption("HR Analytica · 数据分析平台")
        st.caption("Version 1.0.0")


def render_footer() -> None:
    st.markdown("---")
    st.markdown(
        '<div style="text-align:center;color:#94A3B8;font-size:0.82rem;padding:0.5rem 0;">'
        "Powered by 王林 | HR Analytica</div>",
        unsafe_allow_html=True,
    )


def _render_flow() -> None:
    """Three-step flow: upload -> configure -> report."""
    st.markdown('<div class="section-title">三步完成数据分析</div>', unsafe_allow_html=True)
    c1, arrow1, c2, arrow2, c3 = st.columns([1, 0.3, 1, 0.3, 1])
    with c1:
        st.markdown(
            '<div class="step-item"><div class="step-number">1</div>'
            '<div class="step-title">上传 Excel</div>'
            '<div class="step-desc">拖拽上传 CSV 或 Excel</div></div>',
            unsafe_allow_html=True,
        )
    with arrow1:
        st.markdown('<div class="step-arrow">→</div>', unsafe_allow_html=True)
    with c2:
        st.markdown(
            '<div class="step-item"><div class="step-number">2</div>'
            '<div class="step-title">选择分析维度</div>'
            '<div class="step-desc">选择分组列与分析指标</div></div>',
            unsafe_allow_html=True,
        )
    with arrow2:
        st.markdown('<div class="step-arrow">→</div>', unsafe_allow_html=True)
    with c3:
        st.markdown(
            '<div class="step-item"><div class="step-number">3</div>'
            '<div class="step-title">生成专业报告</div>'
            '<div class="step-desc">自动输出图表与结论</div></div>',
            unsafe_allow_html=True,
        )


def render_home() -> None:
    """Landing page."""
    st.markdown(
        '<div class="hero-section">'
        '<div class="hero-title">HR Analytica</div>'
        "</div>",
        unsafe_allow_html=True,
    )

    _render_flow()

    st.markdown('<div class="section-title">六大核心能力</div>', unsafe_allow_html=True)
    features = [
        ("智能解析", "自动识别薪酬、绩效、成本等业务字段", "\U0001f4e7"),
        ("数据清洗", "自动处理缺失值、重复行、混合格式", "\U0001f9f9"),
        ("薪酬诊断", "内置渗透率、公平性、绩效关联分析", "\U0001f50d"),
        ("多维下钻", "支持部门、岗位、渠道交叉分组", "\U0001f4ca"),
        ("可视化", "Plotly交互图表，悬停查看细节", "\U0001f4c8"),
        ("一键报告", "PDF/Excel双格式，专业排版", "\U0001f4cb"),
    ]
    for row in range(2):
        cols = st.columns(3)
        for j in range(3):
            idx = row * 3 + j
            if idx < len(features):
                label, desc, icon = features[idx]
                with cols[j]:
                    st.markdown(
                        f'<div class="feature-card"><div class="feature-icon">{icon}</div>'
                        f'<div class="feature-title">{label}</div><div class="feature-desc">{desc}</div></div>',
                        unsafe_allow_html=True,
                    )

    st.markdown('<div class="section-title">聚焦HR核心分析场景</div>', unsafe_allow_html=True)
    scenes = [
        ("薪酬公平性分析", "上传薪酬数据，一键诊断各部门薪酬渗透率和内部公平性",
         "https://images.unsplash.com/photo-1554224155-8d04cb21cd6c?auto=format&fit=crop&w=600&q=60"),
        ("招聘漏斗优化", "输入招聘各环节数据，自动计算渠道转化率和人均成本",
         "https://images.unsplash.com/photo-1521791136064-7986c2920216?auto=format&fit=crop&w=600&q=60"),
        ("一键专业报告", "自动生成带图表和诊断建议的PDF报告，直接用于工作汇报",
         "https://images.unsplash.com/photo-1454165804606-c3d57bc86b40?auto=format&fit=crop&w=600&q=60"),
    ]
    for i, (title, desc, img) in enumerate(scenes):
        c1, c2 = st.columns([1, 1])
        left, right = (c1, c2) if i % 2 == 0 else (c2, c1)
        with left:
            st.markdown(
                f'<div class="scene-img"><img src="{img}" alt="{title}" '
                'style="width:100%;height:100%;object-fit:cover;border-radius:8px;'
                'filter:grayscale(100%);opacity:0.9"></div>',
                unsafe_allow_html=True,
            )
        with right:
            st.markdown(
                f'<div class="scene-card"><div class="scene-title">{title}</div>'
                f'<div class="scene-desc">{desc}</div></div>',
                unsafe_allow_html=True,
            )



def _plotly_bar(df, gc, vc):
    import plotly.express as px
    gl = gc[0] if isinstance(gc, list) else gc
    agg = df.groupby(gl)[vc].mean().reset_index().sort_values(vc, ascending=False)
    fig = px.bar(agg, x=gl, y=vc, color_discrete_sequence=["#059669"])
    fig.update_layout(plot_bgcolor="white", margin=dict(l=0,r=0,t=0,b=0), height=320)
    fig.update_yaxes(gridcolor="#F2F3F5")
    return fig

def _plotly_hist(df, vc, gc=None):
    import plotly.express as px
    data = df.sample(min(len(df), CONFIG["sample_rows"])) if len(df) > CONFIG["sample_rows"] else df
    gl = gc[0] if isinstance(gc, list) else gc
    if gl and len(data[gl].unique()) <= 10:
        fig = px.histogram(data, x=vc, color=gl, nbins=15, barmode="overlay")
    else:
        fig = px.histogram(data, x=vc, nbins=15, color_discrete_sequence=["#4B5563"])
    fig.update_layout(plot_bgcolor="white", margin=dict(l=0,r=0,t=0,b=0), height=320)
    fig.update_yaxes(gridcolor="#F2F3F5")
    return fig

def _plotly_rank(df, vc, n=15):
    import plotly.express as px
    top = df.nlargest(n, vc).reset_index(drop=True)
    fig = px.bar(top, x=vc, y=top.index + 1, orientation="h",
                 color_discrete_sequence=["#059669"])
    fig.update_layout(plot_bgcolor="white", margin=dict(l=0,r=0,t=0,b=0), height=360,
                      yaxis_title="排名")
    fig.update_xaxes(gridcolor="#F2F3F5")
    return fig


def _render_upload_and_filter(key: str, title: str) -> Optional[pd.DataFrame]:
    """Upload, clean, detect types, preview, and filter. Returns df or None."""
    df = (st.session_state.df_filtered if st.session_state.df_filtered is not None else st.session_state.df)
    if df is not None:
        return df
    st.markdown(f"#### {title}")
    st.markdown("上传文件并配置分析目标后即可自动生成报告。")
    st.markdown(
        '<img class="banner-img" src="https://picsum.photos/id/180/1000/220?grayscale" '
        'alt="数据分析场景">',
        unsafe_allow_html=True,
    )
    up = st.file_uploader("支持 CSV / Excel (.xlsx/.xls)", type=["csv", "xlsx", "xls"], key=f"{key}_up")
    if not up:
        st.info("请上传CSV或Excel文件")
        col1, col2 = st.columns(2)
        bp = BASE_DIR
        if (bp / "薪酬分析样例.csv").exists():
            with open(bp / "薪酬分析样例.csv", "rb") as f:
                col1.download_button(
                    "下载薪酬示例 CSV",
                    data=f.read(),
                    file_name="薪酬分析样例.csv",
                    mime="text/csv",
                )
        if (bp / "薪酬设计全套数据表.xlsx").exists():
            with open(bp / "薪酬设计全套数据表.xlsx", "rb") as f:
                col2.download_button(
                    "下载薪酬示例 Excel",
                    data=f.read(),
                    file_name="薪酬设计全套数据表.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
        return None
    try:
        df_raw, src, sheets = read_file_large(up)
        st.session_state.sheets = sheets
        if sheets:
            sh = st.selectbox("选择 Sheet：", sheets, key=f"{key}_sheet")
            up.seek(0)
            df_raw, src, _ = read_file_large(up, sh)
        df, cr, _, _ = clean_data(df_raw)
        if len(df) > CONFIG["large_threshold"]:
            frac = st.session_state.get("sample_frac", CONFIG["sample_frac"])
            df = df.sample(frac=frac, random_state=42)
            st.info(f"由于数据量较大，已对数据进行{int(frac * 100)}%随机抽样分析，结果具有统计代表性。")
        st.session_state.df = df
        st.session_state.src = src
        st.session_state.df_filtered = None
        with st.expander("数据清洗报告", expanded=True):
            st.write(" - ".join(cr))
        try:
            q = data_quality_report(df)
            st.markdown("**数据质量报告**")
            for c, v in list(q["completeness"].items())[:10]:
                st.progress(float(v), text=f"{c} 完整度 {v * 100:.0f}%")
            st.caption(f"总体评分：{q['grade']}（{q['score']:.2f}）")
            if q["grade"] in ("C", "D"):
                st.warning("数据质量较低，分析结果仅供参考。")
        except Exception:
            pass
        st.markdown(
            '<div style="font-size:1.1rem;font-weight:600;color:#1D2129;">列类型检测</div>',
            unsafe_allow_html=True,
        )
        cat, num, dates, types, idcols = detect_column_types(df)
        st.session_state.cat_cols = cat
        st.session_state.num_cols = num
        st.session_state.date_cols = dates
        st.session_state.col_types = types
        st.session_state.id_cols = idcols
        tags = ""
        for c, t in types.items():
            cls = "tag-numeric" if t == "numeric" else ("tag-date" if t == "date" else "tag-text")
            tags += f'<span class="clean-tag {cls}">{c} ({t})</span> '
        st.markdown(f'<div style="margin:0.5rem 0;">{tags}</div>', unsafe_allow_html=True)
        with st.expander("数据预览", expanded=True):
            st.dataframe(df.head(5), use_container_width=True)
        with st.expander("数据筛选", expanded=False):
            fd = df.copy()
            all_cat = cat + dates
            if all_cat:
                fc = st.selectbox("按列筛选：", all_cat, key=f"{key}_fc")
                fv = st.multiselect(
                    "选择值：",
                    sorted(fd[fc].dropna().unique().astype(str)),
                    key=f"{key}_fv",
                )
                if fv:
                    fd = fd[fd[fc].astype(str).isin(fv)]
            if num:
                fn = st.selectbox("数值范围：", ["不筛选"] + num, key=f"{key}_fn")
                if fn != "不筛选" and fn:
                    mn, mx = float(fd[fn].min()), float(fd[fn].max())
                    rn = st.slider(f"{fn}:", mn, mx, (mn, mx), key=f"{key}_rn")
                    fd = fd[(fd[fn] >= rn[0]) & (fd[fn] <= rn[1])]
            if dates:
                fdt = st.selectbox("日期范围：", ["不筛选"] + dates, key=f"{key}_fdt")
                if fdt != "不筛选" and fdt:
                    dmn, dmx = fd[fdt].min(), fd[fdt].max()
                    dr = st.date_input("日期：", (dmn, dmx), key=f"{key}_dr")
                    if len(dr) == 2:
                        fd = fd[(fd[fdt] >= pd.Timestamp(dr[0])) & (fd[fdt] <= pd.Timestamp(dr[1]))]
            st.info(f"已筛选 {len(df) - len(fd)} 行，剩余 {len(fd)} 行")
            st.session_state.df_filtered = fd
        return fd
    except Exception as e:
        from analysis_engine import logger
        logger.error("Upload error", exc_info=True)
        st.error(f"处理错误: {str(e)}")
        return None


def _render_config(key: str) -> None:
    """Analysis configuration form."""
    c1, c2 = st.columns(2)
    ac = st.session_state.cat_cols + st.session_state.date_cols
    nm = st.session_state.num_cols
    try:
        saved_gc = [c for c in json.loads(load_user_pref("last_group_cols", "[]")) if c in ac]
        saved_vc = load_user_pref("last_value_col", "")
        saved_cc = load_user_pref("last_corr_col", "不分析")
        saved_ex = [c for c in json.loads(load_user_pref("last_extra_charts", "[]"))
                    if c in ["箱线图", "饼图", "热力图", "时间趋势"]]
        default_gc = saved_gc or ac[:1]
        default_vc = saved_vc if saved_vc in nm else (nm[0] if nm else "")
        default_cc = saved_cc if saved_cc in ["不分析"] + nm else "不分析"
    except Exception:
        default_gc = ac[:1]
        default_vc = nm[0] if nm else ""
        default_cc = "不分析"
        saved_ex = []
    with c1:
        gc = st.multiselect("分组列：", ac, default=default_gc, key=f"{key}_gc")
    with c2:
        vc = st.selectbox("分析列：", nm, index=max(nm.index(default_vc), 0) if default_vc in nm else 0, key=f"{key}_vc")
    cc = st.selectbox("关联分析：", ["不分析"] + nm,
                      index=(["不分析"] + nm).index(default_cc) if default_cc in ["不分析"] + nm else 0,
                      key=f"{key}_cc")
    extra = st.multiselect("额外图表：", ["箱线图", "饼图", "热力图", "时间趋势"],
                           default=saved_ex, key=f"{key}_ex")
    if st.button("开始分析", type="primary", use_container_width=True, key=f"{key}_btn"):
        if not gc:
            st.error("请至少选择一个分组列")
        else:
            st.session_state.gc = gc
            st.session_state.vc = vc
            st.session_state.cc = cc
            save_user_pref("last_group_cols", json.dumps(gc, ensure_ascii=False))
            save_user_pref("last_value_col", vc)
            save_user_pref("last_corr_col", cc)
            save_user_pref("last_extra_charts", json.dumps(extra, ensure_ascii=False))
            # extra chart selection is already kept by the widget key
            st.session_state[f"{key}_done"] = True
            st.session_state[f"{key}_step"] = 3
            st.rerun()


def _render_results(key: str, title: str) -> None:
    """Render analysis results with Plotly charts and salary diagnosis."""
    import matplotlib.pyplot as plt
    df = (st.session_state.df_filtered if st.session_state.df_filtered is not None else st.session_state.df)
    gc = st.session_state.gc
    vc = st.session_state.vc
    if isinstance(vc, (list, tuple, np.ndarray)):
        vc = vc[0] if len(vc) else ""
    if not isinstance(vc, str):
        vc = str(vc)
    st.session_state.vc = vc
    cc = st.session_state.cc
    extra = st.session_state.get(f"{key}_ex", [])

    with st.spinner("正在分析数据..."):
        sts = basic_stats(df, gc, vc)
        pen_df = None; fair_r = None; fair_info = None; perf_col = None
        if key == "salary":
            pen_df = calc_penetration(sts, df[vc].mean())
            perf_candidates = find_performance_columns(df)
            with st.sidebar:
                if perf_candidates:
                    perf_col = st.selectbox(
                        "系统检测到以下列可能是绩效/评分数据，请确认用于薪酬诊断的列：",
                        perf_candidates,
                        key=f"{key}_perf",
                    )
                else:
                    perf_col = st.selectbox(
                        "未检测到明确绩效/评分列，请手动指定用于薪酬诊断的列：",
                        ["（不使用）"] + list(df.columns),
                        key=f"{key}_perf",
                    )
            if perf_col and perf_col != "（不使用）":
                fair_r, fair_info = check_internal_fairness(df, vc, perf_col)
        # matplotlib figures for PDF export
        is_money = st.session_state.get("col_types", {}).get(vc, "numeric") != "numeric_non_money"
        bf = plot_bars(sts, money=is_money)
        hf = plot_hist(df, vc, gc[0] if gc else None)
        rf, rd = plot_rank(df, vc, id_cols=st.session_state.get("id_cols", []))
        cf, cv = (plot_corr(df, vc, cc) if cc and cc != "不分析" else (None, None))
        sm = gen_summary(sts, df, gc, vc)
        st.session_state.last_stats = sts
        st.session_state.last_bar_fig = bf
        st.session_state.last_hist_fig = hf
        st.session_state[f"{key}_summary"] = sm

        entry = {"time": datetime.now().strftime("%m-%d %H:%M"),
                 "src": st.session_state.src, "rows": len(df),
                 "group": "+".join(gc), "value": vc,
                 "stats_json": sts.reset_index().to_json(force_ascii=False),
                 "summary": sm}
        h = st.session_state.history; h.insert(0, entry)
        if len(h) > CONFIG["history_limit"]: del h[CONFIG["history_limit"]:]
        st.session_state.history = h
        for fig in [bf, hf, rf, cf]:
            if fig is not None: plt.close(fig)

        level_col = find_level_column(df) if key == "salary" else None
        if isinstance(level_col, (list, tuple)):
            level_col = level_col[0] if level_col else None
        if level_col is not None:
            level_col = str(level_col)
        bw = None
        if level_col:
            bw = salary_bandwidth_analysis(df, level_col, vc)
        try:
            init_db()
            stats_records = sts.reset_index().to_dict(orient="records")
            hid = save_history(
                page_type=key,
                filename=st.session_state.get("src", ""),
                row_count=len(df),
                filter_config={},
                stats_summary=stats_records,
                text_report=sm,
            )
            import io as _io
            for _name, _fig in [("bar", bf), ("hist", hf), ("rank", rf)]:
                if _fig is not None:
                    _buf = _io.BytesIO()
                    _fig.savefig(_buf, format="png", dpi=CONFIG["chart_dpi"])
                    save_report_image(hid, _name, _buf.getvalue())
        except Exception:
            from analysis_engine import logger
            logger.error("History DB save error", exc_info=True)

    ov = df[vc]
    is_money = st.session_state.get("col_types", {}).get(vc, "numeric") != "numeric_non_money"
    def _disp_num(x):
        if x is None or (isinstance(x, float) and pd.isna(x)):
            return "N/A"
        return fmt_val(x) if is_money else f"{x:,.0f}"
    c1, c2, c3, c4 = st.columns(4)
    with c1: stat_card("平均" + vc, _disp_num(ov.mean()), "primary")
    with c2: stat_card("最高" + vc, _disp_num(ov.max()), "success")
    with c3: stat_card("最低" + vc, _disp_num(ov.min()), "warning")
    with c4: stat_card("统计人数", f"{len(ov)} 人")

    if key == "salary":
        st.markdown('<div class="chart-card"><div class="card-title">薪酬诊断</div>', unsafe_allow_html=True)
        if perf_col and perf_col != "（不使用）":
            if fair_r is not None:
                st.markdown(f"- 薪酬与绩效评分整体相关系数：**r={fair_r:.3f}**")
            else:
                st.markdown("- 有效数据不足，无法计算薪酬与绩效的相关系数。")
            if fair_info:
                pc, underpaid, n_top = fair_info
                ratio = underpaid / max(n_top, 1)
                if ratio > 0.3:
                    st.warning(
                        f"⚠️ 存在内部不公平风险：绩效分前25%的员工中，有{underpaid}/{n_top}人的{vc}"
                        "低于整体中位数。建议审视高绩效人群的调薪策略。"
                    )
                else:
                    st.success(
                        f"内部公平性良好：绩效分前25%的员工中，仅{underpaid}/{n_top}人的{vc}"
                        "低于整体中位数。"
                    )
        else:
            st.info("未检测到绩效/评分列，薪资诊断功能暂不可用。如果您确认数据中包含相关列，请在侧边栏手动指定。")
        st.markdown("</div>", unsafe_allow_html=True)

    if key == "salary" and level_col and isinstance(bw, dict) and bw:
        st.markdown('<div class="chart-card"><div class="card-title">薪资带宽分析</div>', unsafe_allow_html=True)
        st.dataframe(bw["table"], use_container_width=True)
        st.plotly_chart(plot_bandwidth_box(df, level_col, vc), use_container_width=True)
        st.plotly_chart(plot_compa_distribution(bw["dist"]), use_container_width=True)
        bw_advice = gen_bandwidth_advice(bw)
        if bw_advice:
            st.markdown(bw_advice)
        st.markdown("</div>", unsafe_allow_html=True)

    col1, col2 = st.columns(2)
    with col1:
        chart_card("各部门平均薪资对比", _plotly_bar(df, gc, vc))
    with col2:
        chart_card("薪资分布直方图", _plotly_hist(df, vc, gc))

    if "箱线图" in extra:
        import plotly.express as px
        gl = gc[0] if isinstance(gc, list) else gc
        chart_card("各部门薪资分布箱线图",
                    px.box(df, x=gl, y=vc, color_discrete_sequence=["#059669"]))
    if "饼图" in extra:
        import plotly.express as px
        gl = gc[0] if isinstance(gc, list) else gc
        cnt = df[gl].value_counts().reset_index()
        gray_palette = ["#111827", "#4B5563", "#6B7280", "#9CA3AF", "#D1D5DB", "#E5E7EB"]
        chart_card("部门人数占比",
                   px.pie(cnt, names=gl, values="count",
                          color_discrete_sequence=gray_palette))
    if "热力图" in extra and len(gc) >= 2:
        pivot = df.pivot_table(index=gc[0], columns=gc[1], values=vc, aggfunc="mean").fillna(0)
        chart_card("交叉平均热力图", plot_heatmap(pivot))
    if "时间趋势" in extra and st.session_state.get("date_cols"):
        date_col = st.session_state.date_cols[0]
        chart_card("时间趋势分析", plot_line(df, date_col, vc, gc[0] if gc else None))

    c1, c2 = st.columns([5, 5])
    with c1:
        st.markdown('<div class="chart-card"><div class="card-title">🏆 薪资排名 Top 15</div></div>',
                    unsafe_allow_html=True)
        st.dataframe(rd.head(15), use_container_width=True, hide_index=True)
    with c2:
        chart_card("薪资排名", _plotly_rank(df, vc))

    st.markdown("---")
    st.markdown('<div class="summary-card"><div class="card-title">智能分析总结</div></div>', unsafe_allow_html=True)
    st.markdown(sm)
    if key == "salary":
        adv = gen_salary_advice(pen_df, fair_r, fair_info, vc)
        if adv:
            st.markdown("### 专业建议")
            st.markdown(adv)

    st.markdown("<br>", unsafe_allow_html=True)
    col_left, _, col_right = st.columns([2, 5, 3])
    with col_left:
        st.markdown('<div class="btn-secondary">', unsafe_allow_html=True)
        if st.button("← 返回重新配置", use_container_width=True, key=f"{key}_reset"):
            st.session_state[f"{key}_done"] = False
            st.session_state.df_filtered = None
            st.rerun()
        st.markdown("</div>", unsafe_allow_html=True)
    with col_right:
        btn1, btn2, btn3 = st.columns(3)
        with btn1:
            st.success("已保存记录")
        with btn2:
            if st.button("导出 Excel", use_container_width=True, key=f"{key}_xlsx"):
                with st.spinner("生成Excel..."):
                    xp = export_excel(df, sts, gc, vc)
                    with open(xp, "rb") as f:
                        st.download_button("下载", data=f.read(),
                                           file_name="data.xlsx",
                                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                           key=f"{key}_dlx")
                    import os
                    try: os.unlink(xp)
                    except Exception: pass
        with btn3:
            if st.button("导出 PDF", use_container_width=True, key=f"{key}_pdf"):
                with st.spinner("生成PDF..."):
                    company_name = load_user_pref("company_name", "HR Analytica")
                    report_title = load_user_pref("report_title", "人力资源数据分析报告")
                    reporter = load_user_pref("reporter", "HR Analytica 自动生成")
                    logo_b64 = load_user_pref("logo_b64", "")
                    logo_path = None
                    if logo_b64:
                        with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as lf:
                            lf.write(base64.b64decode(logo_b64))
                            logo_path = lf.name
                    try:
                        pp = export_pdf(df, sts, gc, vc, bf, hf,
                                        company_name=company_name,
                                        report_title=report_title,
                                        reporter=reporter,
                                        logo_path=logo_path)
                    finally:
                        if logo_path:
                            try: os.unlink(logo_path)
                            except Exception: pass
                    with open(pp, "rb") as f:
                        st.download_button("下载", data=f.read(),
                                           file_name="report.pdf",
                                           mime="application/pdf",
                                           key=f"{key}_dlp")
                    import os
                    try: os.unlink(pp)
                    except Exception: pass


def render_analysis_page(key: str) -> None:
    """Salary/recruit analysis page with three-step navigation."""
    titles = {"salary": "薪酬分析报告", "recruit": "招聘分析报告"}
    title = titles.get(key, "分析报告")
    done_key = f"{key}_done"
    step_key = f"{key}_step"
    if done_key not in st.session_state:
        st.session_state[done_key] = False
    if step_key not in st.session_state:
        st.session_state[step_key] = 1

    st.markdown('<div class="page-breadcrumb">首页 / ' + title + '</div>', unsafe_allow_html=True)
    st.markdown('<div class="page-title">' + title + '</div>', unsafe_allow_html=True)

    step_labels = ["① 数据上传与清洗", "② 分析配置", "③ 分析报告"]
    cols = st.columns(3)
    for col, idx, label in zip(cols, [1, 2, 3], step_labels):
        with col:
            active = st.session_state.get(step_key) == idx
            if st.button(label, key=f"{key}_step_{idx}",
                         type="primary" if active else "secondary",
                         use_container_width=True):
                st.session_state[step_key] = idx
                st.rerun()

    df = (st.session_state.df_filtered if st.session_state.df_filtered is not None else st.session_state.df)
    step = st.session_state.get(step_key, 1)
    if step == 1:
        df2 = _render_upload_and_filter(key, title)
        if df2 is not None:
            st.markdown("<br>", unsafe_allow_html=True)
            if st.button("下一步：分析配置", use_container_width=True, key=f"{key}_next"):
                st.session_state[step_key] = 2
                st.rerun()
    elif step == 2:
        if df is not None:
            _render_config(key)
        else:
            st.info("请先完成数据上传")
    else:
        if df is not None and st.session_state.get(done_key):
            _render_results(key, title)
        else:
            st.info("完成分析配置后，将在此生成报告")


def render_recruit_page() -> None:
    """Dedicated recruitment funnel page."""
    st.markdown('<div class="section-title-inner">招聘效能分析</div>', unsafe_allow_html=True)
    if "recruit_df" not in st.session_state:
        st.session_state.recruit_df = None
    if "recruit_cost_df" not in st.session_state:
        st.session_state.recruit_cost_df = None
    df = st.session_state.recruit_df
    if df is None:
        st.markdown("上传招聘漏斗数据（含应聘渠道/简历数/面试数/面试到场数/录用数/到岗数/招聘成本）。")
        up = st.file_uploader("支持 CSV / Excel", type=["csv", "xlsx", "xls"], key="recruit_up")
        if up:
            try:
                df_raw, src, _ = read_file_large(up)
                if len(df_raw) > CONFIG["large_threshold"]:
                    frac = st.session_state.get("sample_frac", CONFIG["sample_frac"])
                    df_raw = df_raw.sample(frac=frac, random_state=42)
                    st.info(f"由于数据量较大，已对数据进行{int(frac * 100)}%随机抽样分析，结果具有统计代表性。")
                df_raw, cr, _, _ = clean_data(df_raw)
                st.session_state.recruit_df = df_raw
                df = df_raw
                st.session_state.src = src
                with st.expander("数据清洗报告", expanded=True):
                    st.write(" - ".join(cr))
                with st.expander("数据预览", expanded=True):
                    st.dataframe(df_raw.head(10), use_container_width=True)
            except Exception as e:
                from analysis_engine import logger
                logger.error("Recruit upload error", exc_info=True)
                st.error(f"处理错误：{str(e)}")
        else:
            st.info("请上传招聘数据或下载示例文件")
            col1, col2 = st.columns(2)
            bp = BASE_DIR
            if (bp / "招聘分析样例.csv").exists():
                with open(bp / "招聘分析样例.csv", "rb") as f:
                    col1.download_button(
                        "下载招聘示例 CSV",
                        data=f.read(),
                        file_name="招聘分析样例.csv",
                        mime="text/csv",
                    )
            if (bp / "招聘分析样例.xlsx").exists():
                with open(bp / "招聘分析样例.xlsx", "rb") as f:
                    col2.download_button(
                        "下载招聘示例 Excel",
                        data=f.read(),
                        file_name="招聘分析样例.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
            return

    st.markdown("### 渠道成本对照表（可选）")
    cost_up = st.file_uploader(
        "上传成本对照表（需含 应聘渠道 / 渠道总成本 两列）",
        type=["csv", "xlsx", "xls"],
        key="recruit_cost_up",
    )
    if cost_up:
        try:
            cost_raw, _, _ = read_file_large(cost_up)
            cost_raw, _, _, _ = clean_data(cost_raw)
            st.session_state.recruit_cost_df = cost_raw
            st.success("成本对照表已加载")
        except Exception as e:
            from analysis_engine import logger
            logger.error("Cost upload error", exc_info=True)
            st.error(f"成本表处理错误：{str(e)}")

    channel_col = find_channel_column(df) or ("招聘渠道" if "招聘渠道" in df.columns else None)
    if channel_col is None:
        st.error("未检测到渠道列，请确认数据包含 渠道/来源/渠道名称 等字段。")
        return
    detailed = channel_col in df.columns and "简历筛选结果" in df.columns
    required = ["应聘渠道", "简历数", "面试数", "录用数", "到岗数"]
    missing = [x for x in required if x not in df.columns]
    if missing and not detailed:
        st.error(
            f"缺少必要列：{', '.join(missing)}。"
            "请使用含 招聘渠道/简历筛选结果/一面时间/终面结果/是否到岗 的明细数据，"
            "或含 应聘渠道/简历数/面试数/录用数/到岗数 的汇总数据。"
        )
        return
    if "面试到场数" not in df.columns:
        df["面试到场数"] = df.get("面试数", 0)
    if "招聘成本" not in df.columns:
        df["招聘成本"] = 0
    cost_options = ["自动识别", "（不使用）"] + [c for c in df.columns]
    with st.sidebar:
        chosen_cost_col = st.selectbox(
            "招聘成本列（自动识别或手动指定）：",
            cost_options,
            key="recruit_cost_col_sel",
        )
    if chosen_cost_col == "自动识别":
        detail_cost_col = find_cost_column(df)
        if detail_cost_col:
            st.sidebar.caption(f"系统已将【{detail_cost_col}】识别为招聘成本列")
    elif chosen_cost_col == "（不使用）":
        detail_cost_col = None
    else:
        detail_cost_col = chosen_cost_col
    cost_df = st.session_state.get("recruit_cost_df")
    cost_available = cost_df is not None and not cost_df.empty
    if not cost_available:
        cost_available = detail_cost_col is not None
    if not cost_available:
        st.info("未检测到成本数据，请上传渠道成本对照表以计算人均招聘成本。")
    job_col = find_job_column(df)
    cross_mode = "仅渠道"
    if job_col:
        cross_mode = st.selectbox(
            "交叉分析维度：",
            ["仅渠道", "岗位×渠道"],
            key="recruit_cross_mode",
        )
    st.markdown("### 渠道漏斗指标")
    metrics = recruit_metrics(df, cost_df, channel_col)
    st.dataframe(metrics, use_container_width=True)
    csv_download_button(metrics, "渠道漏斗指标.csv", "recruit_metrics_csv")
    st.markdown("### 招聘转化漏斗")
    st.plotly_chart(recruit_funnel(df), use_container_width=True)
    cycle = recruit_cycle_metrics(df, channel_col)
    if cycle:
        st.markdown("### 招聘周期指标")
        if not cycle["cycle_df"].empty:
            st.dataframe(cycle["cycle_df"], use_container_width=True)
        fig_cycle = plot_cycle_bar(cycle["cycle_df"])
        if fig_cycle:
            st.plotly_chart(fig_cycle, use_container_width=True)
        st.markdown(gen_cycle_advice(cycle, df))
    if cross_mode == "岗位×渠道" and job_col:
        import plotly.express as px
        cross = recruit_cross_stats(df, channel_col, job_col)
        st.markdown("### 岗位×渠道交叉统计（录用人数）")
        st.dataframe(cross, use_container_width=True)
        csv_download_button(cross, "岗位渠道交叉统计.csv", "recruit_cross_csv")
        st.plotly_chart(plot_heatmap(cross), use_container_width=True)
        long_df = cross.reset_index().melt(id_vars=job_col, var_name="渠道", value_name="录用人数")
        fig_bar = px.bar(
            long_df,
            x="渠道",
            y="录用人数",
            color=job_col,
            barmode="group",
            color_discrete_sequence=["#059669", "#4B5563", "#111827", "#6B7280", "#9CA3AF"],
        )
        fig_bar.update_layout(plot_bgcolor="white", margin=dict(l=0, r=0, t=40, b=0), height=400)
        st.plotly_chart(fig_bar, use_container_width=True)
    st.markdown("---")
    st.markdown(recruit_summary(metrics, df))
    if st.button("返回重新上传", key="recruit_reset"):
        st.session_state.recruit_df = None
        st.session_state.recruit_cost_df = None
        st.rerun()


def render_report_page() -> None:
    """History and export page backed by SQLite."""
    import json, os
    st.markdown('<div class="section-title-inner">历史报告</div>', unsafe_allow_html=True)
    try:
        init_db()
        history = load_recent_history(10)
    except Exception:
        history = []
    if not history:
        st.info("暂无历史记录，请先完成一次分析")
    else:
        for i, rec in enumerate(history):
            expanded = st.checkbox(
                f"#{i + 1} {rec.get('timestamp', '')} | {rec.get('page_type', '')} | "
                f"{rec.get('filename', '')} | {rec.get('row_count', 0)}行",
                key=f"db_hi_{i}",
            )
            if expanded:
                stats = []
                try:
                    stats = json.loads(rec.get("stats_summary", "[]"))
                    if isinstance(stats, list) and stats:
                        st.dataframe(pd.DataFrame(stats), use_container_width=True)
                except Exception:
                    st.info("无法解析该记录的统计摘要")
                st.markdown(rec.get("text_report", ""))
                for itype, img in get_history_images(rec["id"]):
                    st.image(img, caption=itype, use_container_width=True)
                c1, c2 = st.columns(2)
                with c1:
                    if st.button("重新导出 Excel", key=f"rex_{rec['id']}"):
                        with st.spinner("生成Excel..."):
                            xp = export_history_excel(stats, rec.get("text_report", ""))
                            with open(xp, "rb") as f:
                                st.download_button(
                                    "下载",
                                    data=f.read(),
                                    file_name=f"history_{rec['id']}.xlsx",
                                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                    key=f"rdx_{rec['id']}",
                                )
                            try: os.unlink(xp)
                            except Exception: pass
                with c2:
                    if st.button("重新导出 PDF", key=f"rep_{rec['id']}"):
                        with st.spinner("生成PDF..."):
                            pp = export_history_pdf(stats, rec.get("text_report", ""))
                            with open(pp, "rb") as f:
                                st.download_button(
                                    "下载",
                                    data=f.read(),
                                    file_name=f"history_{rec['id']}.pdf",
                                    mime="application/pdf",
                                    key=f"rdp_{rec['id']}",
                                )
                            try: os.unlink(pp)
                            except Exception: pass
    st.markdown("---")
    if st.session_state.df is not None and st.session_state.last_stats is not None:
        st.markdown("### 导出当前分析")
        df2 = (st.session_state.df_filtered if st.session_state.df_filtered is not None else st.session_state.df)
        s2 = st.session_state.last_stats
        gc2 = st.session_state.gc
        vc2 = st.session_state.vc
        bf2 = st.session_state.last_bar_fig
        hf2 = st.session_state.last_hist_fig
        c1, c2 = st.columns(2)
        with c1:
            if st.button("PDF", type="primary", use_container_width=True, key="ep"):
                with st.spinner("生成PDF..."):
                    pp = export_pdf(df2, s2, gc2, vc2, bf2, hf2)
                    with open(pp, "rb") as f:
                        st.download_button("PDF下载", data=f.read(), file_name="report.pdf", mime="application/pdf", key="dp")
                    try: os.unlink(pp)
                    except Exception: pass
        with c2:
            if st.button("Excel", type="primary", use_container_width=True, key="ee"):
                with st.spinner("生成Excel..."):
                    xp = export_excel(df2, s2, gc2, vc2)
                    with open(xp, "rb") as f:
                        st.download_button("Excel下载", data=f.read(), file_name="data.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="dx")
                    try: os.unlink(xp)
                    except Exception: pass
    else:
        st.info("请先完成分析")
